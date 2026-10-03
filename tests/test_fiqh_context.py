from unittest.mock import Mock, patch

from docx import Document
import pytest

from backend.rag.fiqh.chunker import chunk_record
from backend.rag.fiqh.loader import read_docx
from backend.rag.fiqh.retriever import retrieve_fiqh
from backend.rag.fiqh.schemas import PUBLISHER, SOURCE_NAME
from backend.rag.fiqh.vector_store import FiqhStoreError, QdrantFiqhStore, build_payload, point_id


def source(text):
    return {"source_type": "fiqh", "source_name": SOURCE_NAME, "publisher": PUBLISHER,
            "topic": "اختبار محايد", "volume": None, "page": None, "section": None,
            "text": text, "source_url": "https://example.invalid/test", "verified_source": True}


def test_tiny_leading_and_trailing_segments_are_merged_without_loss():
    text = "عنوان قصير.\n\n" + ("فقرة اصطناعية عن ألوان الرفوف وترتيب الكتب. " * 80) + "\n\nنهاية."
    chunks = chunk_record(source(text), "synthetic.txt")
    assert all(len(c["text"]) >= 100 for c in chunks)
    assert chunks[0]["char_start"] == 0
    assert chunks[-1]["char_end"] == len(text)
    for c in chunks:
        assert c["text"] == text[c["char_start"]:c["char_end"]]
    assert all(b["char_start"] <= a["char_end"] for a, b in zip(chunks, chunks[1:]))


def test_tiny_standalone_document_retained_but_not_indexed():
    chunks = chunk_record(source("عنوان قصير."), "synthetic.txt")
    assert chunks[0]["text"] == "عنوان قصير."
    with pytest.raises(FiqhStoreError, match="100 characters"):
        QdrantFiqhStore().upsert(chunks, [[1.0] * 1024])


def test_actual_heading_structure_is_used_verbatim(tmp_path):
    path = tmp_path / "synthetic.docx"
    document = Document()
    document.add_heading("عنوان اصطناعي محايد", level=1)
    document.add_paragraph("فقرة اصطناعية عن الرفوف. " * 50)
    document.save(path)
    headings = []
    text = read_docx(path, section_spans=headings)
    chunks = chunk_record(source(text), path.name, section_spans=headings)
    assert headings == [(0, "عنوان اصطناعي محايد")]
    assert chunks[0]["section"] == "عنوان اصطناعي محايد"
    assert chunks[0]["text"] == text[chunks[0]["char_start"]:chunks[0]["char_end"]]


def test_neighbors_are_exact_and_never_cross_documents():
    chunks = chunk_record(source("فقرة اصطناعية عن الرفوف والألوان. " * 200), "synthetic.txt")
    other = chunk_record(source("فقرة من وثيقة أخرى. " * 20), "other.txt")
    hit = chunks[1]
    store = Mock()
    store.query.return_value = [{"score": 0.9, "payload": build_payload(hit)}]
    with (
        patch("backend.rag.fiqh.retriever.load_chunks", return_value=other + list(reversed(chunks))),
        patch("backend.rag.fiqh.retriever.QdrantFiqhStore", return_value=store),
        patch("backend.rag.fiqh.retriever.embed_texts", return_value=[[1.0] * 1024]),
    ):
        result = retrieve_fiqh("اختبار")[0]
    assert result["text"] == hit["text"]
    assert result["previous_chunk"]["text"] == chunks[0]["text"]
    assert result["next_chunk"]["text"] == chunks[2]["text"]
    assert result["previous_chunk"]["source"]["input_file"] == "synthetic.txt"


def test_pruning_removes_only_stale_ids_after_upserts():
    chunks = chunk_record(source("فقرة اصطناعية عن الرفوف. " * 20), "synthetic.txt")
    store = QdrantFiqhStore()
    calls = []
    def request(method, path, body):
        calls.append((path, body))
        if path.endswith("scroll"):
            return {"points": [{"id": point_id(chunks[0]["chunk_id"])}, {"id": "obsolete-id"}], "next_page_offset": None}
        return {}
    with patch.object(store, "_request", side_effect=request):
        store.prune_stale_points(chunks)
    assert calls[-1][1] == {"points": ["obsolete-id"]}
    assert {"key": "input_file", "match": {"any": ["synthetic.txt"]}} in calls[0][1]["filter"]["must"]
