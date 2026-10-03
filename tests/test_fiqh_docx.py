"""DOCX fixtures contain synthetic neutral Arabic text only."""
import json

from docx import Document
import pytest
from pydantic import ValidationError

from backend.rag.fiqh.loader import ingest, load_source_file
from backend.rag.fiqh.schemas import PUBLISHER, SOURCE_NAME


def make_docx(tmp_path, paragraphs):
    path = tmp_path / "synthetic.docx"
    document = Document()
    for text in paragraphs:
        document.add_paragraph(text)
    document.save(path)
    return path


def metadata(path, **changes):
    record = {"source_type": "fiqh", "source_name": SOURCE_NAME, "publisher": PUBLISHER,
              "volume": 3, "source_url": "https://example.invalid/synthetic-docx-test",
              "verified_source": True, "topic": "اختبار محايد"}
    record.update(changes)
    path.with_name(path.name + ".metadata.json").write_text(json.dumps(record, ensure_ascii=False), encoding="utf-8")


def test_paragraph_order_exact_arabic_and_only_empty_skipped(tmp_path):
    paragraphs = ["  ألوانُ الرفوفِ: أزرق وأخضر.\u00a0", "", " \t ", "ترتيب الكتب\nعلى الرف.", "", "فقرة أخيرة."]
    path = make_docx(tmp_path, paragraphs)
    metadata(path)
    result = load_source_file(path)[0]
    assert result["text"] == "\n\n".join(text for text in paragraphs if text != "")
    assert result["page"] is None and result["section"] is None


def test_docx_requires_sidecar(tmp_path):
    path = make_docx(tmp_path, ["نص اختبار محايد."])
    with pytest.raises(ValueError, match="Missing provenance sidecar"):
        load_source_file(path)


def test_malformed_sidecar_json_fails(tmp_path):
    path = make_docx(tmp_path, ["نص اختبار محايد."])
    path.with_name(path.name + ".metadata.json").write_text("{invalid", encoding="utf-8")
    with pytest.raises(ValueError, match="Invalid provenance JSON"):
        load_source_file(path)


@pytest.mark.parametrize("changes", [{"source_url": None}, {"publisher": "غير معتمد"}, {"verified_source": False}])
def test_invalid_metadata_does_not_ingest_or_replace_output(tmp_path, changes):
    path = make_docx(tmp_path, ["نص اختبار محايد."])
    metadata(path, **changes)
    output = tmp_path / "processed" / "chunks.json"
    output.parent.mkdir()
    output.write_text("[]", encoding="utf-8")
    with pytest.raises(ValidationError):
        ingest(tmp_path, output)
    assert output.read_text() == "[]"


def test_docx_chunk_integration_and_source_slices(tmp_path):
    paragraphs = [("هذه فقرة اصطناعية عن ترتيب الكتب وألوان الرفوف. " * 30) for _ in range(3)]
    path = make_docx(tmp_path, paragraphs)
    metadata(path, page=12, section="اختبار")
    output = tmp_path / "processed" / "chunks.json"
    count = ingest(tmp_path, output)
    chunks = json.loads(output.read_text(encoding="utf-8"))
    text = "\n\n".join(paragraphs)
    assert count == len(chunks) and count > 1
    for chunk in chunks:
        assert chunk["text"] == text[chunk["char_start"]:chunk["char_end"]]
        assert chunk["input_file"] == "synthetic.docx"
        assert chunk["page"] == 12 and chunk["section"] == "اختبار"


def test_corrupt_docx_fails_clearly(tmp_path):
    path = tmp_path / "corrupt.docx"
    path.write_bytes(b"not a Word document")
    metadata(path)
    with pytest.raises(ValueError, match="Cannot extract DOCX"):
        load_source_file(path)


def test_all_empty_paragraphs_rejected(tmp_path):
    path = make_docx(tmp_path, ["", ""])
    metadata(path)
    with pytest.raises(ValidationError):
        load_source_file(path)


def test_valid_document_discovers_sidecar_and_writes_nonempty_chunks(tmp_path):
    path = make_docx(tmp_path, ["فقرة اصطناعية عن ألوان الرفوف."])
    metadata(path)
    messages = []
    output = tmp_path / "processed" / "chunks.json"
    assert ingest(tmp_path, output, diagnostics=messages.append) == 1
    chunks = json.loads(output.read_text(encoding="utf-8"))
    assert chunks[0]["text"] == "فقرة اصطناعية عن ألوان الرفوف."
    assert f"Metadata sidecar: {path.with_name(path.name + '.metadata.json')}" in messages
    assert any("Found 1" in message for message in messages)
    assert any("Loaded 1 nonempty paragraphs" in message for message in messages)
    assert any("Wrote" in message for message in messages)


def test_tables_nested_and_merged_cells_keep_document_order_without_duplicates(tmp_path):
    path = tmp_path / "synthetic.docx"
    document = Document()
    document.add_paragraph("بداية اصطناعية.")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "خلية أولى."
    table.cell(0, 1).text = "خلية ثانية."
    merged = table.cell(0, 0).merge(table.cell(0, 1))
    nested = merged.add_table(rows=1, cols=1)
    nested.cell(0, 0).text = "خلية متداخلة."
    document.add_paragraph("نهاية اصطناعية.")
    document.save(path)
    metadata(path)
    text = load_source_file(path)[0]["text"]
    assert text == "\n\n".join(["بداية اصطناعية.", "خلية أولى.", "خلية ثانية.", "خلية متداخلة.", "نهاية اصطناعية."])


def test_missing_source_directory_fails(tmp_path):
    with pytest.raises(ValueError, match="Source directory"):
        ingest(tmp_path / "missing", tmp_path / "chunks.json")


def test_cli_invalid_source_exits_nonzero(tmp_path, monkeypatch, capsys):
    from backend.rag.fiqh.loader import main
    make_docx(tmp_path, ["فقرة اصطناعية."])
    monkeypatch.setattr("sys.argv", ["loader", "--raw-dir", str(tmp_path), "--out", str(tmp_path / "chunks.json")])
    with pytest.raises(SystemExit) as error:
        main()
    assert error.value.code == 1
    assert "Missing provenance sidecar" in capsys.readouterr().err
